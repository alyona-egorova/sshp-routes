"""Regenerate routes.js from a filled workbook. Usage: python make_routes_js.py SSHP-routesDB-filled.xlsx
Reads columns Name, Strava Link, Distance, Elevation, Direction from the first sheet (any extra columns ignored)."""
import json, sys
from openpyxl import load_workbook
grp = {'W':'West','SW':'West','NW':'North-West','N':'North','NE':'North-East','E':'East','SE':'East','S':'South','Local':'Local'}
ws = load_workbook(sys.argv[1], data_only=True).active
hdr = [str(c.value).strip().lower() if c.value else '' for c in ws[1]]
col = lambda *k: next(i for i,h in enumerate(hdr) if any(x in h for x in k))
iN, iL, iD, iE, iR = col('name'), col('link','strava'), col('dist'), col('elev'), col('direc')
out = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[iL] or row[iD] in (None, ''): continue
    out.append({"n": str(row[iN]).strip(), "u": row[iL], "k": round(float(row[iD]),1), "e": int(row[iE]), "d": grp.get(row[iR], row[iR])})
open('routes.js','w', encoding='utf-8').write('window.ROUTES = ' + json.dumps(out, ensure_ascii=False, indent=0) + ';\n')
print(f'wrote routes.js with {len(out)} routes')
